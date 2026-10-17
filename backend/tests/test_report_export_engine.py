import io
import sys
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.services.report_export_engine import dashboard_to_excel, dashboard_to_pdf


def test_dashboard_excel_export_keeps_untrusted_strings_literal():
    payload = dashboard_to_excel(
        kpis={"=1+1": "=HYPERLINK(\"https://example.invalid\",\"click\")"},
        monthly=[{"month": "  =1+1", "revenue": 100}],
        expense_breakdown_rows=[{"category": "+SUM(1,1)", "amount": 10}],
    )
    workbook = load_workbook(io.BytesIO(payload), data_only=False)

    assert workbook["KPI Summary"]["A2"].value == "'=1+1"
    assert workbook["KPI Summary"]["B2"].value.startswith("'=HYPERLINK")
    assert workbook["Monthly Trend"]["A2"].value == "'  =1+1"
    assert workbook["Expense Breakdown"]["A2"].value == "'+SUM(1,1)"
    assert not [
        cell.coordinate
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.data_type == "f"
    ]


def test_dashboard_pdf_escapes_untrusted_paragraph_markup():
    kpis = {
        "revenue": 100,
        "expenses": 40,
        "net_profit": 60,
        "profit_margin_pct": 60,
        "cash_position": 100,
        "receivables": 0,
        "payables": 0,
        "gst_liability": 0,
        "tds_payable": 0,
        "working_capital": 100,
    }
    payload = dashboard_to_pdf(
        org_name="<b>Acme & Co</b>",
        period_label="<period>",
        kpis=kpis,
        monthly=[],
    )
    extracted = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(payload)).pages)
    assert "<b>Acme & Co</b>" in extracted
    assert "<period>" in extracted
