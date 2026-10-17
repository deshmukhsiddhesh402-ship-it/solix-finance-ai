import io

import pandas as pd
from openpyxl import load_workbook

from app.routers.excel_automation import _safe_excel_dataframe


def test_cleaned_export_neutralizes_formula_like_cells_and_headers():
    formula_header = '=HYPERLINK("https://example.invalid","click")'
    original = pd.DataFrame({formula_header: [
        "=1+1",
        "  =2+2",
        chr(160) + "=3+3",
        "+SUM(1,1)",
        "-1+2",
        "@SUM(A1:A2)",
        chr(9) + "=1+1",
        "safe text",
    ]})

    safe = _safe_excel_dataframe(original)
    buffer = io.BytesIO()
    safe.to_excel(buffer, index=False, engine="openpyxl")
    workbook = load_workbook(io.BytesIO(buffer.getvalue()), data_only=False)
    sheet = workbook.active

    assert sheet["A1"].value.startswith("'=")
    assert [sheet.cell(row=row, column=1).value for row in range(2, 8)] == [
        "'=1+1", "'  =2+2", "'" + chr(160) + "=3+3", "'+SUM(1,1)", "'-1+2", "'@SUM(A1:A2)"
    ]
    assert sheet["A8"].value == "'" + chr(9) + "=1+1"
    assert sheet["A9"].value == "safe text"
    assert not [cell.coordinate for row in sheet.iter_rows() for cell in row if cell.data_type == "f"]
    assert original.columns[0] == formula_header


def test_safe_export_handles_duplicate_headers_after_sanitization():
    original = pd.DataFrame([["=1+1", "=2+2"]], columns=["=header", "'=header"])
    safe = _safe_excel_dataframe(original)
    assert list(safe.columns) == ["'=header", "'=header"]
    assert list(safe.iloc[0]) == ["'=1+1", "'=2+2"]
    assert list(original.columns) == ["=header", "'=header"]
